import io

import openpyxl

from mailpipe.chunk import split_text
from mailpipe.clean import Boilerplate, clean_body, strip_quoted
from mailpipe.extract import MixedWordSplitter, extract_cpu, extract_dates, windows
from mailpipe.lang import detect_language, estimate_tokens, has_simplified, to_zh_tw, trim_to_tokens
from mailpipe.parse_xlsx import parse_xlsx
from mailpipe.thread import normalize_subject

PATTERNS = {
    "po_number": r"\bPO[-\s]?\d{6,10}\b",
    "amount": r"(?:NT\$|NTD|USD|US\$|新台幣)\s?[\d,]+(?:\.\d+)?(?:\s?元)?",
}


def test_strip_quoted_zh_outlook_header():
    body = "好的，收到。\n\n寄件者： 王小明\n寄件日期： 2025年10月6日\n收件者： Bob\n主旨： RE: 報價\n\n舊的內容"
    assert strip_quoted(body) == "好的，收到。"


def test_strip_quoted_original_message_and_wrote():
    assert strip_quoted("hi\n-----Original Message-----\nFrom: x") == "hi"
    assert strip_quoted("thanks\n\nOn Mon, Oct 6, 2025 at 2:30 PM Bob wrote:\n> old") == "thanks"
    assert strip_quoted("ok\n\n王小明 於 2025年10月6日 寫道：\n> 舊") == "ok"


def test_strip_quoted_keeps_inline_quotes_with_new_text():
    body = "see below\n> question 1\nanswer 1\n> question 2\nanswer 2"
    assert strip_quoted(body) == body


def test_boilerplate_signature_removed():
    sig = "王小明\n業務部 經理"
    bp = Boilerplate(global_min=3, sender_min=2, disclaimers=[])
    bodies = [f"內容 {i}\n\n{sig}" for i in range(3)]
    for b in bodies:
        bp.observe("ming@x", b)
    assert clean_body(bodies[0], "ming@x", bp) == "內容 0"


def test_dates_roc_zh_and_relative():
    sent = "2025-10-06T10:00:00+08:00"  # a Monday
    got = {e.value: e.normalized for e in extract_dates("交期 115/01/15，請於10月15日前回覆，下週五開會", sent)}
    assert got["115/01/15"] == "2026-01-15"
    assert got["10月15日"] == "2025-10-15"
    assert got["下週五"] == "2025-10-17"
    en = {e.value.lower(): e.normalized for e in extract_dates("by next Friday or October 10, 2025", sent)}
    assert en["next friday"] == "2025-10-17"
    assert en["october 10, 2025"] == "2025-10-10"


def test_dates_month_rollover():
    got = {e.value: e.normalized for e in extract_dates("1月5日交貨", "2025-12-20T09:00:00+08:00")}
    assert got["1月5日"] == "2026-01-05"


def test_regex_entities_fullwidth():
    ents = {(e.type, e.normalized) for e in extract_cpu("ＰＯ－１２３４５６７８ 金額 NT$ 3,400,000", None, PATTERNS)}
    assert ("amount", "TWD 3400000") in ents


def test_xlsx_header_detection_and_merged_cells():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Title"
    ws.merge_cells("A1:C1")
    ws.append([])
    ws.append(["PO", "Qty", "Status"])
    ws.append(["PO-1", 5, "ok"])
    ws.append(["PO-2", 7, None])
    ws["C4"] = "merged"
    ws.merge_cells("C4:C5")
    h = wb.create_sheet("hidden")
    h.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_xlsx(buf.getvalue())
    assert parsed.hidden_skipped == ["hidden"]
    s = parsed.sheets[0]
    assert s.columns == ["PO", "Qty", "Status"]
    assert s.rows == [["PO-1", "5", "merged"], ["PO-2", "7", "merged"]]


def test_language_and_opencc():
    assert detect_language("請確認報價單內容並回覆") == "zh"
    assert detect_language("Please confirm the quotation") == "en"
    assert detect_language("請確認 quotation for AB-12345 and the delivery schedule") == "mixed"
    assert to_zh_tw("请确认软件信息") == "請確認軟體資訊"
    assert has_simplified("请确认") and not has_simplified("請確認")


def test_token_trim():
    t = "中" * 1000
    assert estimate_tokens(trim_to_tokens(t, 100)) <= 110


def test_normalize_subject():
    assert normalize_subject("RE: FW: 回覆: 報價") == normalize_subject("報價")
    assert normalize_subject("ＲＥ：報價") == "報價"


def test_split_and_windows():
    parts = split_text("段落。" * 400, 200)
    assert all(estimate_tokens(p) <= 200 for p in parts) and len(parts) > 1
    sp = MixedWordSplitter()
    toks = [t for t, _, _ in sp("請 Contoso 提供 AB-12345 報價")]
    assert "Contoso" in toks and "AB-12345" in toks
    w = windows("請提供報價單 AB-12345 " * 200, sp, size=300, overlap=30)
    assert len(w) >= 3 and w[0][0] == 0


def test_regex_no_double_tagging_and_boundaries():
    from mailpipe.config import load_config

    cfg = load_config()
    ents = extract_cpu("PO-20251001 USD 9,250, tracking 1Z999AA10123456784 call 02-2345-6789", None, cfg.regex)
    got = {(e.type, e.value) for e in ents}
    assert ("po_number", "PO-20251001") in got and ("part_number", "PO-20251001") not in got
    assert ("amount", "USD 9,250") in got
    assert ("phone", "02-2345-6789") in got
    assert not any(t == "phone" and "0123456784" in v for t, v in got)


def test_html_inline_tags_keep_sentences():
    from mailpipe.parse_msg import html_to_text

    text = html_to_text("<p>Let's meet <b>this Thursday</b> 2pm.</p><div>Second<br>line</div>")
    assert "Let's meet this Thursday 2pm." in text
    assert "Second\nline" in text


def test_normalize_due():
    from mailpipe.categorize import normalize_due

    assert normalize_due("2025-10-17", None) == "2025-10-17"
    assert normalize_due("下週五", "2025-10-06T10:00:00+08:00") == "2025-10-17"
    assert normalize_due("as soon as possible", None) is None
