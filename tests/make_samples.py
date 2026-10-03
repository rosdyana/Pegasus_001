"""Generate anonymized fixtures in tests/samples (run: python tests/make_samples.py)."""
from __future__ import annotations

import io
import sys
from pathlib import Path

import docx
import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from msgwriter import Attachment, Message  # noqa: E402

OUT = Path(__file__).parent / "samples"

SIG_WANG = "王小明 Ming Wang\n業務部 經理\n範例科技股份有限公司\nTel: 02-2345-6789 #123"
DISCLAIMER = ("本郵件可能包含機密資訊，僅供指定收件人使用。若您非指定收件人，請立即通知寄件人並刪除本郵件。"
              "This e-mail may contain confidential information intended only for the addressee.")
CONV = bytes.fromhex("01d8f0a1b2c3d4e5f60718293a4b5c6d7e8f9012ab")


def make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("出貨檢驗報告", level=1)
    d.add_paragraph("料號 AB-12345 於 114/10/02 完成出貨檢驗。")
    d.add_heading("檢驗結果", level=2)
    d.add_paragraph("外觀檢查合格", style="List Bullet")
    d.add_paragraph("尺寸量測 3 件超出公差", style="List Bullet")
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate([["項目", "數量", "結果"], ["外觀", "100", "OK"], ["尺寸", "100", "NG 3"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_heading("後續處理", level=2)
    d.add_paragraph("請供應商於 10月15日 前提出改善對策 (8D report)。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_quote_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "報價單"
    ws["A1"] = "範例科技 報價單"
    ws.merge_cells("A1:D1")
    ws["A2"] = "日期: 2025-10-06"
    ws.append([])
    ws.append(["料號", "品名", "數量", "單價 (USD)"])
    for i in range(1, 46):
        ws.append([f"AB-{12300 + i}", f"連接器 type {i}", i * 10, round(1.5 + i * 0.1, 2)])
    hidden = wb.create_sheet("internal_cost")
    hidden.append(["料號", "成本"])
    hidden.append(["AB-12301", 0.9])
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_standalone_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly shipments"
    ws.append(["Week 40 shipment summary"])
    ws.append([])
    ws.append(["PO", "Customer", "Ship date", "Qty", "Status"])
    rows = [["PO-20251001", "Contoso Ltd", "2025-10-01", 500, "shipped"],
            ["PO-20251002", "Fabrikam Inc", "2025-10-03", 1200, "customs hold"],
            ["PO-20251003", "台灣範例電子", "2025-10-04", 300, "shipped"]]
    for r in rows:
        ws.append(r)
    ws["F3"] = "Remark"
    ws.merge_cells("F4:F6")
    ws["F4"] = "see email"
    wb.save(path)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    msgs = {
        "01_rfq.msg": Message(
            subject="詢價：連接器 AB-12345 報價需求", sender_name="Bob Chen", sender_email="bob.chen@contoso.example",
            to="王小明 <ming.wang@example.com.tw>", date="Thu, 02 Oct 2025 09:15:00 +0800",
            message_id="<rfq-001@contoso.example>", conversation_index=CONV,
            body=("王經理您好，\n\n我們需要 AB-12345 連接器 5,000 pcs，請於下週五前提供報價，"
                  "交期希望在 115/01/15 之前。\n\n謝謝\n\nBob Chen\nPurchasing, Contoso Ltd\n\n" + DISCLAIMER),
        ),
        "02_quote_reply.msg": Message(
            subject="RE: 詢價：連接器 AB-12345 報價需求", sender_name="王小明", sender_email="ming.wang@example.com.tw",
            to="Bob Chen <bob.chen@contoso.example>", cc="Amy Lin <amy.lin@example.com.tw>",
            date="Mon, 06 Oct 2025 14:30:00 +0800", message_id="<quote-002@example.com.tw>",
            in_reply_to="<rfq-001@contoso.example>", conversation_index=CONV + bytes.fromhex("0000a1b2c3"),
            body=("Bob 您好，\n\n附件為報價單，單價 USD 1.85，總金額 USD 9,250，報價有效期限至 2025/10/31。\n"
                  "如確認請回覆 PO 號碼。\n\n" + SIG_WANG + "\n\n" + DISCLAIMER + "\n\n"
                  "-----原始郵件-----\n寄件者: Bob Chen <bob.chen@contoso.example>\n寄件日期: 2025年10月2日 上午 09:15\n"
                  "收件者: 王小明\n主旨: 詢價：連接器 AB-12345 報價需求\n\n王經理您好，我們需要 AB-12345 連接器 5,000 pcs"),
            attachments=[Attachment("報價單_AB-12345.xlsx", make_quote_xlsx(),
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")],
        ),
        "03_po.msg": Message(
            subject="RE: 詢價：連接器 AB-12345 報價需求", sender_name="Bob Chen", sender_email="bob.chen@contoso.example",
            to="王小明 <ming.wang@example.com.tw>", date="Wed, 08 Oct 2025 10:00:00 +0800",
            message_id="<po-003@contoso.example>", in_reply_to="<quote-002@example.com.tw>",
            conversation_index=CONV + bytes.fromhex("0000a1b2c3") + bytes.fromhex("0000d4e5f6"),
            body=("Hi Ming,\n\nPlease find our PO-20251008 for 5,000 pcs AB-12345 at USD 1.85. "
                  "Please confirm delivery by next Friday.\n\nThanks,\nBob\n\n" + DISCLAIMER + "\n\n"
                  "On Mon, Oct 6, 2025 at 2:30 PM 王小明 wrote:\n> Bob 您好，附件為報價單"),
        ),
        "04_quality_issue.msg": Message(
            subject="【品質異常】AB-12345 尺寸超出公差", sender_name="Amy Lin", sender_email="amy.lin@example.com.tw",
            to="QA Team <qa@example.com.tw>", cc="王小明 <ming.wang@example.com.tw>",
            date="Fri, 03 Oct 2025 16:45:00 +0800", message_id="<qa-004@example.com.tw>",
            body=("各位好，\n\n出貨檢驗發現 AB-12345 有 3 件尺寸超出公差，詳見附件檢驗報告。"
                  "請 QA 於 10月15日 前提出 8D report，並評估是否需要 RMA。\n\nAmy\n--\nAmy Lin 林小美\n品保部"),
            attachments=[Attachment("出貨檢驗報告.docx", make_docx(),
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                         Attachment("photo.pdf", b"%PDF-1.4 fake", "application/pdf")],
        ),
        "05_forward_nested.msg": Message(
            subject="FW: Shipment delay notice", sender_name="王小明", sender_email="ming.wang@example.com.tw",
            to="Amy Lin <amy.lin@example.com.tw>", date="Tue, 07 Oct 2025 11:20:00 +0800",
            message_id="<fw-005@example.com.tw>",
            body="Amy，\n\n轉寄 Fabrikam 的延遲通知，請協助確認海關狀態。\n\n" + SIG_WANG + "\n\n" + DISCLAIMER,
            attachments=[Attachment("Shipment delay notice.msg", embedded=Message(
                subject="Shipment delay notice", sender_name="Logistics", sender_email="logistics@fabrikam.example",
                to="ming.wang@example.com.tw", date="Mon, 06 Oct 2025 08:00:00 +0000",
                message_id="<delay-100@fabrikam.example>",
                body=("Dear customer,\n\nShipment for PO-20251002 is on customs hold in Taipei. "
                      "Expected release: October 10, 2025. Tracking: 1Z999AA10123456784.\n\nFabrikam Logistics"),
            ))],
        ),
        "06_meeting_html.msg": Message(
            subject="Weekly sync: Q4 planning", sender_name="Amy Lin", sender_email="amy.lin@example.com.tw",
            to="Team <team@example.com.tw>", date="Mon, 06 Oct 2025 09:00:00 +0800",
            message_id="<mtg-006@example.com.tw>",
            html=("<html><head><meta charset='utf-8'><style>p{color:red}</style></head><body>"
                  "<p>Hi all,</p><p>Let's meet <b>this Thursday</b> 2pm in room 3A to review Q4 forecast. "
                  "請 Ming 準備 Contoso 的銷售數字。</p><p>Amy</p></body></html>"),
        ),
        "07_approval.msg": Message(
            subject="請核准：Contoso 特殊價格申請", sender_name="王小明", sender_email="ming.wang@example.com.tw",
            to="Director Huang <huang@example.com.tw>", date="Thu, 09 Oct 2025 13:00:00 +0800",
            message_id="<appr-007@example.com.tw>",
            body=("黃總您好，\n\nContoso 要求 AB-12345 單價降至 USD 1.75 (原 USD 1.85)，"
                  "年用量約 60,000 pcs，金額 NT$ 3,400,000。請於明天下班前核准。\n\n" + SIG_WANG + "\n\n" + DISCLAIMER),
        ),
    }
    for name, m in msgs.items():
        (OUT / name).write_bytes(m.to_bytes())
    make_standalone_xlsx(OUT / "08_weekly_shipments.xlsx")
    print(f"wrote {len(msgs) + 1} fixtures to {OUT}")


if __name__ == "__main__":
    main()
