"""End-to-end runs against the sample fixtures with a fake Ollama server and a fake GLiNER."""
import json
import sqlite3

import pytest
from openpyxl import load_workbook

from fake_ollama import FakeOllama
from mailpipe.config import load_config
from mailpipe.lang import has_simplified
from mailpipe.pipeline import Pipeline

MODELS = ["qwen3:14b", "qwen3:8b", "bge-m3:latest"]


def make_cfg(url, samples, out, **llm):
    cfg = load_config(None, input_dir=str(samples), output_dir=str(out))
    cfg.llm.base_url = url
    cfg.discovery.min_emails = 3
    for k, v in llm.items():
        setattr(cfg.llm, k, v)
    return cfg


@pytest.fixture
def ollama():
    with FakeOllama(MODELS) as f:
        yield f


def db(out):
    c = sqlite3.connect(out / "mailpipe.db")
    c.row_factory = sqlite3.Row
    return c


def test_full_run_then_rerun(ollama, samples, tmp_path, fake_gliner):
    out = tmp_path / "out"
    cfg = make_cfg(ollama.url, samples, out)
    stats = Pipeline(cfg, cpu_only=True).run()
    c = db(out)

    files = {r["status"]: r["n"] for r in c.execute("SELECT status, COUNT(*) n FROM files GROUP BY status")}
    assert files == {"parsed": 8}
    # 7 top-level emails + 1 nested
    assert c.execute("SELECT COUNT(*) FROM emails").fetchone()[0] == 8
    assert c.execute("SELECT COUNT(*) FROM emails WHERE parent_email_id IS NOT NULL").fetchone()[0] == 1
    # every .msg has a non-empty body (html-only one included)
    assert c.execute("SELECT COUNT(*) FROM emails WHERE length(body_raw)=0").fetchone()[0] == 0

    # RFQ / quote / PO are one thread
    tids = {r[0] for r in c.execute("SELECT thread_id FROM emails WHERE subject LIKE '%AB-12345 報價需求%'")}
    assert len(tids) == 1

    # the quote reply lost its quoted history, signature and disclaimer
    body = c.execute("SELECT body_clean FROM emails WHERE message_id='<quote-002@example.com.tw>'").fetchone()[0]
    assert "原始郵件" not in body and "機密" not in body and "USD 1.85" in body

    # attachments: xlsx with hidden sheet skipped, docx parsed, pdf unsupported
    sheets = [r["name"] for r in c.execute("SELECT name FROM sheets")]
    assert "報價單" in sheets and "internal_cost" not in sheets
    assert c.execute("SELECT status FROM attachments WHERE kind='pdf'").fetchone()[0] == "unsupported"
    docx_md = c.execute("SELECT text_md FROM attachments WHERE kind='docx'").fetchone()[0]
    assert "# 出貨檢驗報告" in docx_md and "| 項目 | 數量 | 結果 |" in docx_md

    ents = {(r["type"], r["normalized"] or r["value"]) for r in c.execute("SELECT * FROM entities")}
    assert ("po_number", "PO20251008") in ents
    assert ("date", "2026-01-15") in ents          # ROC 115/01/15
    assert ("company", "Contoso") in ents          # GLiNER (fake)

    # LLM outputs: labels for all emails, zh-TW only after OpenCC
    assert c.execute("SELECT COUNT(*) FROM labels").fetchone()[0] == 8
    for (text,) in c.execute("SELECT summary FROM threads WHERE summary IS NOT NULL UNION ALL "
                             "SELECT action FROM action_items"):
        assert not has_simplified(text), text
    assert c.execute("SELECT SUM(changed) FROM opencc_log").fetchone()[0] > 0

    # chunks, FTS (trigram matches Chinese substrings and part numbers), vectors
    n_chunks = c.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    assert n_chunks > 8
    assert c.execute("SELECT COUNT(*) FROM chunks WHERE embedded=0").fetchone()[0] == 0
    hit = c.execute("SELECT COUNT(*) FROM chunks_fts WHERE chunks_fts MATCH '\"AB-12345\"'").fetchone()[0]
    assert hit > 0
    assert c.execute("SELECT COUNT(*) FROM chunks_fts WHERE chunks_fts MATCH '\"尺寸超出\"'").fetchone()[0] > 0
    lines = (out / "chunks.jsonl").read_text(encoding="utf-8").splitlines()
    assert len(lines) == n_chunks and json.loads(lines[0])["chunk_id"]

    for f in ("report.html", "report.md", "topic_discovery.md", "preflight.json"):
        if f != "preflight.json":
            assert (out / f).exists(), f
    html = (out / "report.html").read_text(encoding="utf-8")
    assert "Category distribution" in html and "Quality samples" in html
    assert stats["llm"]["calls"] > 0

    # Ollama models are unloaded after each GPU phase
    assert not ollama.state.loaded

    # re-run: nothing new, no LLM or embedding calls
    before = len(ollama.state.calls)
    stats2 = Pipeline(cfg, cpu_only=True).run()
    chat_calls = [x for x in ollama.state.calls[before:] if x["path"] in ("/api/chat", "/api/embed")
                  and x.get("keep_alive") != 0]
    assert chat_calls == []
    assert stats2["discover"]["skipped_done"] == 8 and stats2["discover"]["new"] == 0
    assert stats2["total_s"] < 10


def test_oom_falls_back_to_smaller_model(ollama, samples, tmp_path, fake_gliner):
    ollama.state.oom_models.add("qwen3:14b")
    out = tmp_path / "out"
    cfg = make_cfg(ollama.url, samples, out)
    stats = Pipeline(cfg, cpu_only=True, limit=5).run()
    c = db(out)
    rows = c.execute("SELECT model, fallback_used, method FROM labels").fetchall()
    assert rows and all(r["model"] == "qwen3:8b" and r["fallback_used"] == 1 for r in rows)
    assert stats["llm"]["oom_events"] >= 3
    # after 3 straight fallbacks the phase sticks to the smaller model
    assert any(r["method"] == "llm:primary" for r in rows)


def test_oom_everywhere_marks_failed_and_continues(ollama, samples, tmp_path, fake_gliner):
    ollama.state.oom_models.update({"qwen3:14b", "qwen3:8b"})
    out = tmp_path / "out"
    cfg = make_cfg(ollama.url, samples, out)
    Pipeline(cfg, cpu_only=True, limit=2).run()
    c = db(out)
    statuses = {r[0] for r in c.execute("SELECT status FROM progress WHERE stage='classify'")}
    assert statuses == {"failed_oom"}
    assert "failed_oom" in (out / "report.html").read_text(encoding="utf-8")


def test_limit_then_resume(ollama, samples, tmp_path, fake_gliner):
    out = tmp_path / "out"
    cfg = make_cfg(ollama.url, samples, out)
    s1 = Pipeline(cfg, cpu_only=True, limit=3).run()
    assert s1["discover"]["new"] == 3
    s2 = Pipeline(cfg, cpu_only=True).run()
    assert s2["discover"]["new"] == 5 and s2["discover"]["skipped_done"] == 3


def test_review_roundtrip(ollama, samples, tmp_path, fake_gliner):
    from mailpipe.review import compute_metrics, export_review, import_review
    from mailpipe.store import Store

    out = tmp_path / "out"
    cfg = make_cfg(ollama.url, samples, out)
    Pipeline(cfg, cpu_only=True).run()
    store = Store(out / "mailpipe.db")
    path = export_review(store, cfg, out / "review")
    wb = load_workbook(path)
    assert wb["Categories"].column_dimensions["A"].hidden
    ws = wb["Categories"]
    header = [c.value for c in ws[1]]
    ok_col = header.index("correct? (Y/N)") + 1
    corr_col = header.index("corrected_category") + 1
    miss_col = header.index("missing entities") + 1
    ws.cell(2, ok_col, "Y")
    ws.cell(3, ok_col, "N")
    ws.cell(3, corr_col, "other")
    ws.cell(2, miss_col, "Contoso Ltd, 5,000 pcs")
    ws = wb["Entities"]
    h = [c.value for c in ws[1]]
    ws.cell(2, h.index("correct? (Y/N)") + 1, "Y")
    ws.cell(3, h.index("correct? (Y/N)") + 1, "N")
    ws = wb["Summaries"]
    h = [c.value for c in ws[1]]
    ws.cell(2, h.index("rating (1-3)") + 1, 3)
    wb["Info"]["B1"] = "Specialist A"
    wb.save(path)

    res = import_review(store, cfg, path)
    assert res["reviewer"] == "Specialist A"
    assert res["counts"]["category"] == 2
    m = compute_metrics(store, cfg)
    assert m["category_n"] == 2
    assert 0 <= m["category_accuracy"] <= 1
    assert m["missing_count"] == 2
    assert m["summary_rating"] == 3
    # predictions are untouched
    assert store.scalar("SELECT COUNT(*) FROM labels") == 8
    # re-import is idempotent
    import_review(store, cfg, path)
    assert store.scalar("SELECT COUNT(*) FROM reviews WHERE item_type='category'") == 2


def test_search(ollama, samples, tmp_path, fake_gliner):
    from mailpipe.embed import Embedder, search
    from mailpipe.store import Store

    out = tmp_path / "out"
    cfg = make_cfg(ollama.url, samples, out)
    Pipeline(cfg, cpu_only=True).run()
    store = Store(out / "mailpipe.db")
    hits = search(store, out / "vectors", Embedder(cfg), "PO-20251002 customs", k=5)
    assert any("PO-20251002" in h["text"] for h in hits)
    hits = search(store, out / "vectors", Embedder(cfg), "尺寸", k=3)   # < 3 chars: LIKE fallback
    assert any("尺寸" in h["text"] for h in hits)


def test_corrupt_and_duplicate_files_do_not_stop_the_batch(ollama, samples, tmp_path, fake_gliner):
    (samples / "broken.msg").write_bytes(b"not an ole file")
    sub = samples / "sub"
    sub.mkdir()
    (sub / "copy_of_rfq.msg").write_bytes((samples / "01_rfq.msg").read_bytes())
    out = tmp_path / "out"
    cfg = make_cfg(ollama.url, samples, out)
    stats = Pipeline(cfg, cpu_only=True).run()
    c = db(out)
    row = c.execute("SELECT status, error FROM files WHERE path LIKE '%broken.msg'").fetchone()
    assert row["status"] == "failed" and row["error"]
    # nothing half-written for the failed file
    fid = c.execute("SELECT file_id FROM files WHERE path LIKE '%broken.msg'").fetchone()[0]
    assert c.execute("SELECT COUNT(*) FROM emails WHERE file_id=?", (fid,)).fetchone()[0] == 0
    assert stats["discover"]["duplicates"] == 1
    assert c.execute("SELECT COUNT(*) FROM files WHERE status='parsed'").fetchone()[0] == 8
    assert "broken.msg" in (out / "report.html").read_text(encoding="utf-8")
