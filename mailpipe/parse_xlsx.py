"""Stage 2: .xlsx to tables (Parquet) + Markdown preview."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
import pandas as pd

HEADER_SCAN_ROWS = 20
PREVIEW_ROWS = 20


@dataclass
class Sheet:
    name: str
    header_row: int            # 1-based row number in the sheet
    columns: list[str]
    rows: list[list[str | None]]
    quality_flag: str | None = None

    @property
    def n_rows(self) -> int:
        return len(self.rows)

    @property
    def n_cols(self) -> int:
        return len(self.columns)

    def dataframe(self) -> pd.DataFrame:
        return pd.DataFrame(self.rows, columns=self.columns, dtype="string")

    def preview_md(self, n: int = PREVIEW_ROWS) -> str:
        return md_table(self.columns, self.rows[:n])


@dataclass
class Workbook:
    sheets: list[Sheet] = field(default_factory=list)
    hidden_skipped: list[str] = field(default_factory=list)


def cell_text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat(sep=" ", timespec="minutes") if (v.hour or v.minute) else v.date().isoformat()
    if isinstance(v, (date, time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s or None


def md_escape(s: str | None) -> str:
    return (s or "").replace("|", "\\|").replace("\n", " ")


def md_table(columns: list[str], rows: list[list[str | None]]) -> str:
    if not columns:
        return ""
    lines = ["| " + " | ".join(md_escape(c) for c in columns) + " |",
             "| " + " | ".join("---" for _ in columns) + " |"]
    lines += ["| " + " | ".join(md_escape(v) for v in r) + " |" for r in rows]
    return "\n".join(lines)


def detect_header(grid: list[list[str | None]]) -> int:
    """0-based index of the row (within the first 20) with the most distinct non-empty text cells."""
    best, best_score = 0, -1
    for i, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        # distinct values, so a merged title spread over the row does not win
        score = len({v for v in row if v is not None and not _is_number(v)})
        if score > best_score:
            best, best_score = i, score
    return best


def _is_number(s: str) -> bool:
    return bool(re.fullmatch(r"[-+]?[\d,]*\.?\d+%?", s))


def _unique(names: list[str | None]) -> list[str]:
    out, seen = [], {}
    for i, n in enumerate(names):
        base = (n or f"col_{i + 1}").strip() or f"col_{i + 1}"
        k = seen.get(base, 0)
        seen[base] = k + 1
        out.append(base if k == 0 else f"{base}_{k + 1}")
    return out


def read_sheet(ws) -> Sheet | None:
    # fill merged cells down and right with the top-left value
    merged_values = {}
    for rng in list(ws.merged_cells.ranges):
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged_values[(r, c)] = v

    grid: list[list[str | None]] = []
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        grid.append([cell_text(merged_values.get((r_idx, c_idx), v)) for c_idx, v in enumerate(row, start=1)])

    # trim empty trailing rows / columns
    while grid and all(v is None for v in grid[-1]):
        grid.pop()
    if not grid:
        return None
    width = max((max((i + 1 for i, v in enumerate(r) if v is not None), default=0) for r in grid), default=0)
    if width == 0:
        return None
    grid = [(r + [None] * width)[:width] for r in grid]
    # drop leading fully-empty columns
    lead = min((next((i for i, v in enumerate(r) if v is not None), width) for r in grid))
    grid = [r[lead:] for r in grid]

    h = detect_header(grid)
    columns = _unique(grid[h])
    body = [r for r in grid[h + 1:] if any(v is not None for v in r)]

    flags = []
    filled = sum(1 for v in grid[h] if v is not None) / max(len(columns), 1)
    if filled < 0.5:
        flags.append("sparse_header")
    gaps = _empty_row_gaps(grid[h + 1:])
    if gaps:
        flags.append("multiple_blocks")
    if not body:
        flags.append("no_data_rows")
    return Sheet(name=ws.title, header_row=h + 1, columns=columns, rows=body,
                 quality_flag=",".join(flags) or None)


def _empty_row_gaps(rows: list[list[str | None]]) -> int:
    """Count runs of 2+ empty rows between data: a hint that the sheet holds several tables."""
    gaps, run, seen_data = 0, 0, False
    for r in rows:
        if any(v is not None for v in r):
            if seen_data and run >= 2:
                gaps += 1
            seen_data, run = True, 0
        else:
            run += 1
    return gaps


def parse_xlsx(data: bytes) -> Workbook:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    out = Workbook()
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                out.hidden_skipped.append(ws.title)
                continue
            sheet = read_sheet(ws)
            if sheet is not None:
                out.sheets.append(sheet)
    finally:
        wb.close()
    return out


def safe_name(s: str) -> str:
    return re.sub(r"[^\w.-]+", "_", s, flags=re.UNICODE)[:60] or "sheet"


def write_parquet(sheet: Sheet, tables_dir: Path, sha: str, index: int) -> Path:
    tables_dir.mkdir(parents=True, exist_ok=True)
    path = tables_dir / f"{sha[:16]}_{index}_{safe_name(sheet.name)}.parquet"
    sheet.dataframe().to_parquet(path, index=False)
    return path


def workbook_markdown(wb: Workbook) -> str:
    parts = []
    for s in wb.sheets:
        parts.append(f"## Sheet: {s.name} ({s.n_rows} rows x {s.n_cols} cols)\n\n{s.preview_md()}")
    return "\n\n".join(parts)
