"""Stage 9: specialist review workbook export/import and accuracy metrics."""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .store import Store, now

HIGHLIGHT = PatternFill("solid", fgColor="FFF2CC")
HEADER = PatternFill("solid", fgColor="D9E1F2")
INPUT = PatternFill("solid", fgColor="E2EFDA")
YES = {"y", "yes", "1", "true", "是", "對", "对", "正確", "o", "v"}
NO = {"n", "no", "0", "false", "否", "錯", "错", "不正確", "x"}

# sheet -> (item_type, columns); input columns are marked with a trailing '*'
SHEETS = {
    "Categories": ("category", ["id", "review first", "subject", "sender", "date", "body preview", "predicted category",
                                "confidence", "needs_reply", "priority", "model", "correct? (Y/N)*",
                                "corrected_category*", "missing entities*", "comment*"]),
    "Entities": ("entity", ["id", "email subject", "type", "value", "context", "method", "confidence",
                            "correct? (Y/N)*", "corrected_value*", "comment*"]),
    "Action items": ("action", ["id", "email subject", "action", "owner", "due_date", "model",
                                "valid? (Y/N)*", "corrected_text*", "comment*"]),
    "Summaries": ("summary", ["id", "topic", "emails", "summary", "model", "rating (1-3)*", "comment*"]),
}
WIDTHS = {"subject": 40, "email subject": 30, "body preview": 60, "context": 50, "summary": 80, "action": 50,
          "comment*": 30, "missing entities*": 30, "corrected_category*": 22, "corrected_text*": 40,
          "corrected_value*": 22, "topic": 40, "sender": 26, "value": 24}


def _yn(v) -> int | None:
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in YES:
        return 1
    if s in NO:
        return 0
    return None


def _prev_reviews(store: Store, item_type: str) -> dict[str, dict]:
    return {r["item_id"]: dict(r) for r in store.q("SELECT * FROM reviews WHERE item_type=?", (item_type,))}


def _yn_text(v) -> str | None:
    return None if v is None else ("Y" if v else "N")


def export_review(store: Store, cfg, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"review_{date.today().isoformat()}.xlsx"
    n = 2
    while path.exists():  # never overwrite a workbook the specialist may be filling in
        path = out_dir / f"review_{date.today().isoformat()}_{n}.xlsx"
        n += 1
    wb = Workbook()
    info = wb.active
    info.title = "Info"
    info["A1"], info["B1"] = "Reviewer name:", ""
    info["B1"].fill = INPUT
    info["A3"] = "How to review"
    info["A3"].font = Font(bold=True)
    for i, line in enumerate([
        "Fill only the green columns. Yellow rows are the ones to review first "
        "(low confidence, category 'other', or produced by a fallback model).",
        "correct? / valid?: Y or N. If N, put the right value in the corrected column.",
        "Categories: use the dropdown for corrected_category. Note entities the pipeline missed in 'missing entities', "
        "separated by commas.",
        "Summaries: rate 1 (wrong or useless), 2 (partly right), 3 (good).",
        "You can review in several sessions. Every import keeps the latest answer per row.",
        "Do not edit or delete the hidden 'id' column.",
    ], start=4):
        info[f"A{i}"] = line
    info.column_dimensions["A"].width = 120

    lists = wb.create_sheet("Lists")
    for i, name in enumerate(cfg.category_names, start=1):
        lists[f"A{i}"] = name
    lists.sheet_state = "hidden"

    th = cfg.thresholds.review_confidence
    sheets = {name: wb.create_sheet(name) for name in SHEETS}
    for name, ws in sheets.items():
        cols = SHEETS[name][1]
        ws.append([c.rstrip("*") for c in cols])
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(1, i)
            cell.font = Font(bold=True)
            cell.fill = INPUT if c.endswith("*") else HEADER
            ws.column_dimensions[cell.column_letter].width = WIDTHS.get(c, 14)
        ws.column_dimensions["A"].hidden = True
        ws.freeze_panes = "C2"

    # Categories
    prev = _prev_reviews(store, "category")
    prev_missing = _prev_reviews(store, "missing_entities")
    rows = []
    for r in store.q("SELECT e.email_id, e.subject, e.sender, e.sent_at, e.body_clean, l.category, l.confidence, "
                     "l.needs_reply, l.priority, l.model, l.fallback_used FROM emails e JOIN labels l USING(email_id)"):
        hi = (r["confidence"] or 0) < th or r["category"] == "other" or r["fallback_used"]
        p = prev.get(str(r["email_id"]), {})
        rows.append((not hi, r["confidence"] or 0, [
            r["email_id"], "yes" if hi else "", r["subject"], r["sender"], (r["sent_at"] or "")[:16],
            (r["body_clean"] or "")[:300], r["category"], round(r["confidence"] or 0, 2),
            "yes" if r["needs_reply"] else "no", r["priority"], r["model"], _yn_text(p.get("is_correct")),
            p.get("corrected_value"), prev_missing.get(str(r["email_id"]), {}).get("corrected_value"), p.get("comment"),
        ], hi))
    _fill(sheets["Categories"], rows)

    # Entities
    prev = _prev_reviews(store, "entity")
    rows = []
    for r in store.q("SELECT n.entity_id, e.subject, n.type, n.value, n.context, n.method, n.confidence "
                     "FROM entities n LEFT JOIN emails e USING(email_id) ORDER BY n.email_id, n.type"):
        hi = (r["confidence"] or 1) < th
        p = prev.get(str(r["entity_id"]), {})
        rows.append((not hi, 0, [r["entity_id"], r["subject"], r["type"], r["value"], r["context"], r["method"],
                                 round(r["confidence"] or 0, 2), _yn_text(p.get("is_correct")),
                                 p.get("corrected_value"), p.get("comment")], hi))
    _fill(sheets["Entities"], rows, sort=False)

    # Action items
    prev = _prev_reviews(store, "action")
    rows = []
    for r in store.q("SELECT a.action_id, e.subject, a.action, a.owner, a.due_date, a.model, l.fallback_used "
                     "FROM action_items a LEFT JOIN emails e ON e.email_id=a.email_id "
                     "LEFT JOIN labels l ON l.email_id=a.email_id ORDER BY a.email_id"):
        hi = bool(r["fallback_used"])
        p = prev.get(str(r["action_id"]), {})
        rows.append((not hi, 0, [r["action_id"], r["subject"], r["action"], r["owner"], r["due_date"], r["model"],
                                 _yn_text(p.get("is_correct")), p.get("corrected_value"), p.get("comment")], hi))
    _fill(sheets["Action items"], rows)

    # Summaries
    prev = _prev_reviews(store, "summary")
    rows = []
    for r in store.q("SELECT thread_id, topic, email_count, summary, model FROM threads WHERE summary IS NOT NULL "
                     "ORDER BY email_count DESC"):
        hi = bool(r["model"]) and r["model"] != cfg.llm.model
        p = prev.get(r["thread_id"], {})
        rows.append((not hi, 0, [r["thread_id"], r["topic"], r["email_count"], r["summary"], r["model"],
                                 p.get("rating"), p.get("comment")], hi))
    _fill(sheets["Summaries"], rows)

    _validations(sheets, len(cfg.category_names))
    wb.save(path)
    return path


def _fill(ws, rows: list[tuple], sort: bool = True) -> None:
    if sort:
        rows.sort(key=lambda r: (r[0], r[1]))
    else:
        rows.sort(key=lambda r: r[0])
    for _, _, values, hi in rows:
        ws.append(values)
        if hi:
            for cell in ws[ws.max_row]:
                cell.fill = HIGHLIGHT
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _col(ws, header: str) -> str | None:
    for cell in ws[1]:
        if cell.value == header:
            return cell.column_letter
    return None


def _validations(sheets: dict, n_categories: int) -> None:
    for name, ws in sheets.items():
        n = max(ws.max_row, 2)
        for header in ("correct? (Y/N)", "valid? (Y/N)"):
            col = _col(ws, header)
            if col:
                dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{col}2:{col}{n}")
        col = _col(ws, "corrected_category")
        if col:
            dv = DataValidation(type="list", formula1=f"=Lists!$A$1:$A${n_categories}", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{n}")
        col = _col(ws, "rating (1-3)")
        if col:
            dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{n}")


# --------------------------------------------------------------------------- import

def import_review(store: Store, cfg, path: Path, reviewer: str | None = None) -> dict:
    wb = load_workbook(path, data_only=True)
    if reviewer is None and "Info" in wb.sheetnames:
        reviewer = (wb["Info"]["B1"].value or "").strip() or None
    reviewer = reviewer or "unknown"
    counts: dict[str, int] = defaultdict(int)
    valid_cats = set(cfg.category_names)
    warnings = []
    for sheet, (item_type, cols) in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(header) if h}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[idx["id"]] is None:
                continue
            item_id = str(row[idx["id"]])
            get = lambda h: (row[idx[h]] if h in idx and idx[h] < len(row) else None)  # noqa: E731
            comment = get("comment")
            if item_type == "category":
                ok, corr = _yn(get("correct? (Y/N)")), get("corrected_category")
                pred = get("predicted category")
                if corr and corr not in valid_cats:
                    warnings.append(f"email {item_id}: unknown category '{corr}'")
                if corr and corr != pred:
                    ok = 0
                elif corr and corr == pred and ok is None:
                    ok = 1
                missing = get("missing entities")
                if missing:
                    _save(store, "missing_entities", item_id, None, str(missing), None, None, reviewer)
                    counts["missing_entities"] += 1
                if ok is None and not comment:
                    continue
                _save(store, item_type, item_id, ok, corr, None, comment, reviewer)
            elif item_type in ("entity", "action"):
                ok_col = "correct? (Y/N)" if item_type == "entity" else "valid? (Y/N)"
                corr = get("corrected_value" if item_type == "entity" else "corrected_text")
                ok = _yn(get(ok_col))
                if ok is None and corr:
                    ok = 0
                if ok is None and not comment:
                    continue
                _save(store, item_type, item_id, ok, corr, None, comment, reviewer)
            else:
                rating = get("rating (1-3)")
                try:
                    rating = int(rating) if rating not in (None, "") else None
                except (TypeError, ValueError):
                    warnings.append(f"thread {item_id}: rating '{rating}' is not 1-3")
                    rating = None
                if rating is None and not comment:
                    continue
                _save(store, item_type, item_id, None, None, rating, comment, reviewer)
            counts[item_type] += 1
    store.commit()
    return {"reviewer": reviewer, "counts": dict(counts), "warnings": warnings}


def _save(store: Store, item_type, item_id, ok, corrected, rating, comment, reviewer) -> None:
    store.upsert("reviews", {
        "item_type": item_type, "item_id": item_id, "is_correct": ok,
        "corrected_value": None if corrected in (None, "") else str(corrected), "rating": rating,
        "comment": None if comment in (None, "") else str(comment), "reviewer": reviewer, "reviewed_at": now(),
    }, key=("item_type", "item_id"))


# --------------------------------------------------------------------------- metrics

def compute_metrics(store: Store, cfg) -> dict:
    m: dict = {"reviewed": store.scalar("SELECT COUNT(*) FROM reviews") or 0}
    if not m["reviewed"]:
        return m
    # categories
    pairs = []
    for r in store.q("SELECT r.is_correct, r.corrected_value, l.category FROM reviews r "
                     "JOIN labels l ON l.email_id = CAST(r.item_id AS INTEGER) WHERE r.item_type='category'"):
        if r["corrected_value"]:
            truth = r["corrected_value"]
        elif r["is_correct"] == 1:
            truth = r["category"]
        else:
            continue  # marked wrong without a correction: counts as wrong, truth unknown
        pairs.append((r["category"], truth))
    wrong_unknown = store.scalar("SELECT COUNT(*) FROM reviews WHERE item_type='category' AND is_correct=0 "
                                 "AND corrected_value IS NULL") or 0
    n = len(pairs) + wrong_unknown
    m["category_n"] = n
    m["category_accuracy"] = (sum(p == t for p, t in pairs) / n) if n else None
    per = defaultdict(lambda: [0, 0])
    confusion: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for p, t in pairs:
        per[p][0] += p == t
        per[p][1] += 1
        confusion[p][t] += 1
    m["per_category"] = {c: {"n": v[1], "acc": v[0] / v[1]} for c, v in sorted(per.items())}
    m["confusion"] = {k: dict(v) for k, v in confusion.items()}
    m["confusion_labels"] = sorted({x for pair in pairs for x in pair})
    # entities
    ent = defaultdict(lambda: [0, 0])
    tp = 0
    for r in store.q("SELECT n.type, r.is_correct FROM reviews r JOIN entities n "
                     "ON n.entity_id = CAST(r.item_id AS INTEGER) WHERE r.item_type='entity' AND r.is_correct IS NOT NULL"):
        ent[r["type"]][0] += r["is_correct"]
        ent[r["type"]][1] += 1
        tp += r["is_correct"]
    m["entity_precision"] = {t: {"n": v[1], "precision": v[0] / v[1]} for t, v in sorted(ent.items())}
    missing = 0
    for r in store.q("SELECT corrected_value FROM reviews WHERE item_type='missing_entities'"):
        # split on separators, but not on thousands commas like "5,000"
        missing += len([x for x in re.split(r"(?<!\d),|,(?!\d)|[，、;；\n]", r[0] or "") if x.strip()])
    m["missing_count"] = missing
    m["entity_recall"] = tp / (tp + missing) if (tp + missing) else None
    # actions + summaries
    a = store.one("SELECT COUNT(*) n, SUM(is_correct) ok FROM reviews WHERE item_type='action' AND is_correct IS NOT NULL")
    m["action_n"] = a["n"]
    m["action_validity"] = (a["ok"] / a["n"]) if a["n"] else None
    m["summary_rating"] = store.scalar("SELECT AVG(rating) FROM reviews WHERE item_type='summary' AND rating IS NOT NULL")
    return m
